#!/usr/bin/env python3
"""Build SourceData.xlsx from panel-level CSV files.

The figure scripts write one CSV file per figure panel. This helper combines
those CSV files into a single Excel workbook with one sheet per CSV. Each sheet
starts with a short title/description in row 1, followed by the data table.
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def safe_sheet_name(stem: str, used: set[str]) -> str:
    base = re.sub(r"[^A-Za-z0-9_]", "_", stem)[:31] or "Sheet"
    sheet = base
    i = 1
    while sheet in used:
        suffix = f"_{i}"
        sheet = (base[: 31 - len(suffix)] + suffix)[:31]
        i += 1
    used.add(sheet)
    return sheet


META_COLUMNS = ["source_figure", "source_panel", "source_description"]


def set_reasonable_widths(ws, rows: list[list[str]], max_scan_rows: int = 200) -> None:
    if not rows:
        return
    widths = [0] * max(len(row) for row in rows)
    for row in rows[:max_scan_rows]:
        for idx, value in enumerate(row):
            widths[idx] = max(widths[idx], min(len(str(value)), 60))
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = max(10, width + 2)


def split_metadata(rows: list[list[str]], file_stem: str) -> tuple[str, str, str, list[list[str]]]:
    """Return figure/panel/description plus rows without repeated metadata columns."""
    if not rows:
        return "", "", file_stem, rows

    header = rows[0]
    meta_present = header[: len(META_COLUMNS)] == META_COLUMNS
    if not meta_present:
        return "", "", file_stem, rows

    data_rows = [row[len(META_COLUMNS) :] for row in rows]
    first_data = rows[1] if len(rows) > 1 else []
    source_figure = first_data[0] if len(first_data) > 0 else ""
    source_panel = first_data[1] if len(first_data) > 1 else ""
    source_description = first_data[2] if len(first_data) > 2 else file_stem
    return source_figure, source_panel, source_description, data_rows


def format_data_sheet(ws, rows: list[list[str]], title: str) -> None:
    max_cols = max((len(row) for row in rows), default=1)
    if max_cols > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.font = Font(bold=True, size=11)
    title_cell.alignment = Alignment(wrap_text=True, vertical="top")

    if rows:
        for cell in ws[3]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(max_cols)}{max(ws.max_row, 3)}"


def build_workbook(panel_dir: Path, output_xlsx: Path) -> None:
    files = sorted(panel_dir.glob("*.csv"))
    if not files:
        raise FileNotFoundError(f"No CSV files found in {panel_dir}")

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)
    used = set()

    for file in files:
        sheet = safe_sheet_name(file.stem, used)
        ws = wb.create_sheet(sheet)
        raw_rows: list[list[str]] = []
        with file.open(newline="", encoding="utf-8") as handle:
            reader = csv.reader(handle)
            for row in reader:
                raw_rows.append(row)

        source_figure, source_panel, source_description, rows = split_metadata(raw_rows, file.stem)
        title_parts = []
        if source_figure:
            title_parts.append(source_figure)
        if source_panel:
            title_parts.append(f"panel {source_panel}")
        prefix = " ".join(title_parts)
        title = f"{prefix}: {source_description}" if prefix else source_description

        ws.append([title])
        ws.append([])
        for row in rows:
            ws.append(row)
        format_data_sheet(ws, rows, title)
        set_reasonable_widths(ws, [[title]] + rows)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    if output_xlsx.exists():
        output_xlsx.unlink()
    wb.save(output_xlsx)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--panel-dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build_workbook(args.panel_dir, args.output)


if __name__ == "__main__":
    main()
