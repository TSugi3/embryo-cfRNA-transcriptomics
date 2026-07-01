#!/usr/bin/env python3
"""Build revised Extended Data Figure 3 as a multi-slide working PowerPoint file."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


from path_config import BASE_DIR, FIG_DIR
OUT_PPTX = FIG_DIR / "ExtendedDataFigure3_Revised_working.pptx"

SLIDE_W = 7.5
SLIDE_H = 10.833333333333334


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.31, 0.24, 9, bold=True)


def add_picture(slide, filename: str, left: float, top: float, width: float, height: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide1, "Extended Data Figure 3 (continues)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture(slide1, "FigureS3A_Distribution_KS_Density.png", 0.208, 0.420, 7.083, 1.771)
    add_picture(slide1, "FigureS3B_Reactome_UpregulatedOnly.png", 0.775, 2.455, 5.950, 8.429)
    add_panel_label(slide1, "A", 0.208, 0.278)
    add_panel_label(slide1, "B", 0.208, 2.310)

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide2, "Extended Data Figure 3 (continued)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture(slide2, "FigureS3C_Reactome_DownregulatedOnly.png", 0.775, 0.560, 5.950, 8.429)
    add_panel_label(slide2, "C", 0.208, 0.420)

    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide3, "Extended Data Figure 3 (continued)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture(slide3, "FigureS3D_GSEA_Reactome_dotplot.png", 0.208, 0.420, 7.083, 3.229)
    add_picture(slide3, "FigureS3E_GSEAplot_EWEvsAWE_Reactome.png", 0.208, 4.100, 7.083, 5.508)
    add_panel_label(slide3, "D", 0.208, 0.278)
    add_panel_label(slide3, "E", 0.208, 3.958)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
