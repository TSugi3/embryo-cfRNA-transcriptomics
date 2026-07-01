#!/usr/bin/env python3
"""Build revised Figure 5 as a two-slide working PowerPoint file."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


from path_config import BASE_DIR, FIG_DIR
OUT_PPTX = FIG_DIR / "Figure5_Revised_working.pptx"

SLIDE_W = 7.5
SLIDE_H = 10.833333333333334


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.31, 0.24, 9, bold=True)


def add_picture(slide, filename: str, left: float, top: float, width: float, height: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def add_picture_keep_aspect(slide, filename: str, left: float, top: float, width: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide1, "Figure 5 (continues)", 0.0, 0.005, 1.99, 0.30, 10)
    add_picture(slide1, "Figure5A_UpSetPlot_nonDEG_CV.png", 0.230, 0.459, 1.958, 1.958)
    add_picture(slide1, "Figure5B_Biotype_PieChart.png", 2.195, 0.459, 5.111, 1.958)
    add_picture_keep_aspect(slide1, "Figure5C_GO_BP_dotplot_viridis.png", 0.208, 2.442, 7.083)
    add_picture_keep_aspect(slide1, "Figure5D_heatmap_autophagy.png", 0.223, 5.300, 7.083)
    add_panel_label(slide1, "A", 0.208, 0.296)
    add_panel_label(slide1, "B", 2.195, 0.296)
    add_panel_label(slide1, "C", 0.208, 2.271)
    add_panel_label(slide1, "D", 0.208, 5.182)

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide2, "Figure 5 (continued)", 0.0, 0.005, 1.99, 0.30, 10)
    add_picture(slide2, "Figure5E_chord_autophagy.png", 0.208, 0.308, 7.083, 6.500)
    add_panel_label(slide2, "E", 0.208, 0.296)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
