#!/usr/bin/env python3
"""Build final Figure 3 as an editable PowerPoint slide."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


from path_config import FIG_DIR

OUT_PPTX = FIG_DIR / "Figure3_editable.pptx"
SLIDE_W = 7.5
SLIDE_H = 8.3


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.31, 0.24, 9, bold=True)


def add_picture(slide, filename: str, left: float, top: float, width: float, height: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def add_picture_keep_aspect(slide, filename: str, left: float, top: float, width: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, "Figure 3", 0.0, 0.005, 1.25, 0.30, 10)
    add_picture_keep_aspect(slide, "Figure3A_DEG_Proportion.png", 0.208, 0.365, 7.083)
    add_picture_keep_aspect(slide, "Figure3B_GO_BP_UpregulatedOnly.png", 0.208, 1.545, 7.083)
    add_picture_keep_aspect(slide, "Figure3C_GO_BP_DownregulatedOnly.png", 0.208, 4.985, 7.083)
    add_panel_label(slide, "A", 0.208, 0.296)
    add_panel_label(slide, "B", 0.208, 1.425)
    add_panel_label(slide, "C", 0.208, 4.865)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
