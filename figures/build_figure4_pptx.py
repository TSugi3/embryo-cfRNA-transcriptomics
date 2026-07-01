#!/usr/bin/env python3
"""Build revised Figure 4 as a working PowerPoint slide."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


from path_config import BASE_DIR, FIG_DIR
OUT_PPTX = FIG_DIR / "Figure4_Revised_working.pptx"

SLIDE_W = 7.5
SLIDE_H = 10.833333333333334


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.31, 0.24, 9, bold=True)


def add_picture(slide, filename: str, left: float, top: float, width: float, height: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, "Figure 4", 0.0, 0.005, 1.25, 0.30, 10)

    add_picture(slide, "Figure4A_GSEA_GO_BP_ESMvsASM_dotplot.png", 0.208, 0.308, 7.083, 3.931)
    add_picture(slide, "Figure4B_Representative_Gene_Expression_Boxplot.png", 0.208, 4.238, 7.083, 1.569)
    add_picture(slide, "Figure4C_ExosomeMarkers_Boxplot.png", 0.208, 6.007, 7.083, 1.486)
    add_picture(slide, "Figure4D_ssGSEA_Boxplot_GO_BP.png", 0.208, 7.596, 7.083, 1.569)

    add_panel_label(slide, "A", 0.208, 0.296)
    add_panel_label(slide, "B", 0.208, 4.131)
    add_panel_label(slide, "C", 0.208, 5.889)
    add_panel_label(slide, "D", 0.208, 7.540)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
