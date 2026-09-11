#!/usr/bin/env python3
"""Small helpers for assembling editable figure PowerPoint files."""

from __future__ import annotations

from typing import Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = 7.5


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.31, 0.24, 9, bold=True)


def add_picture(slide, fig_dir, filename: str, left: float, top: float,
                width: float, height: Optional[float] = None) -> None:
    path = fig_dir / filename
    if not path.exists():
        raise FileNotFoundError(path)
    if height is None:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    else:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                                 width=Inches(width), height=Inches(height))


def build_single_slide(out_pptx, title: str, fig_dir, slide_h: float,
                       pictures: list[tuple], labels: list[tuple[str, float, float]]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(slide_h)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, title, 0.0, 0.005, 3.4, 0.30, 10)

    for spec in pictures:
        if len(spec) == 5:
            filename, left, top, width, height = spec
            add_picture(slide, fig_dir, filename, left, top, width, height)
        elif len(spec) == 4:
            filename, left, top, width = spec
            add_picture(slide, fig_dir, filename, left, top, width)
        else:
            raise ValueError(f"Invalid picture specification: {spec!r}")

    for label, left, top in labels:
        add_panel_label(slide, label, left, top)

    prs.save(str(out_pptx))
    print(out_pptx)
