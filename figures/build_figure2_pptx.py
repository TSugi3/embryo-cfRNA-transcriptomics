#!/usr/bin/env python3
"""Build Figure 2 as an editable PowerPoint slide."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


from path_config import BASE_DIR, FIG_DIR
OUT_PPTX = FIG_DIR / "Figure2_editable.pptx"

SLIDE_W = 7.5
SLIDE_H = 8.32


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False, align=PP_ALIGN.LEFT) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.20, 0.18, 8, bold=True)


def add_rect(slide, text: str, left: float, top: float, width: float, height: float,
             fill: str = "FFFFFF", line: str = "1F2F50", font_size: float = 7.0,
             bold: bool = False) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.color.rgb = RGBColor.from_string(line)
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_down_arrow(slide, x: float, y: float) -> None:
    add_textbox(slide, "↓", x - 0.08, y, 0.16, 0.16, 9.5, bold=True, align=PP_ALIGN.CENTER)


def add_figure2a(slide) -> None:
    panel = FIG_DIR / "Figure2A_from_source_PPT.png"
    if panel.exists():
        add_picture(slide, panel, 0.62, 0.42, 6.26, 1.99)
    else:
        add_rect(slide, "Schematic workflow panel prepared separately",
                 0.62, 0.42, 6.26, 1.99, fill="F7F7F7", line="D0D0D0", font_size=8.0)


def add_picture(slide, path: Path, left: float, top: float, width: float, height: float) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_textbox(slide, "Figure 2", 0.00, 0.00, 1.25, 0.30, 10)
    add_figure2a(slide)

    add_picture(slide, FIG_DIR / "Figure2B_GO_BP_EWEvsESM_dotplot.png", 0.30, 2.58, 4.32, 2.32)
    add_picture(slide, FIG_DIR / "Figure2C_UpSetPlot_nonDEG.png", 4.63, 2.58, 2.75, 2.32)
    add_picture(slide, FIG_DIR / "Figure2D_ssGSEA_heatmap.png", 0.31, 4.98, 7.03, 1.42)
    add_picture(slide, FIG_DIR / "Figure2E_ssGSEA_score_comparison.png", 0.21, 6.48, 7.08, 1.77)

    add_panel_label(slide, "A", 0.21, 0.30)
    add_panel_label(slide, "B", 0.21, 2.40)
    add_panel_label(slide, "C", 4.54, 2.40)
    add_panel_label(slide, "D", 0.21, 4.81)
    add_panel_label(slide, "E", 0.21, 6.31)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
