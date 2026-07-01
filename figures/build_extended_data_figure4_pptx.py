#!/usr/bin/env python3
"""Build Extended Data Figure 4 as a editable PowerPoint file."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


from path_config import BASE_DIR, FIG_DIR
OUT_PPTX = FIG_DIR / "ExtendedDataFigure4_editable.pptx"

SLIDE_W = 7.5
SLIDE_H = 10.833333333333334


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.31, 0.24, 9, bold=True)


def add_picture(slide, filename: str, left: float, top: float, width: float, height: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def add_picture_keep_aspect(slide, filename: str, left: float, top: float, width: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide1, "Extended Data Figure 4 (continues)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture(slide1, "FigureS4A_KEGG_dotplot.png", 0.208, 0.420, 7.083, 3.149)
    add_picture(slide1, "FigureS4B_heatmap_vesicle.png", 0.208, 3.805, 7.083, 4.724)
    add_panel_label(slide1, "A", 0.208, 0.278)
    add_panel_label(slide1, "B", 0.208, 3.660)

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide2, "Extended Data Figure 4 (continued)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture_keep_aspect(slide2, "FigureS4C_chord_vesicle.png", 0.208, 0.420, 7.083)
    add_panel_label(slide2, "C", 0.208, 0.278)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
