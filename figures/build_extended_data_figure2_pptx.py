#!/usr/bin/env python3
"""Build Extended Data Figure 2 as a two-slide editable PowerPoint file."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


from path_config import BASE_DIR, FIG_DIR
OUT_PPTX = FIG_DIR / "ExtendedDataFigure2_editable.pptx"

SLIDE_W = 7.5
SLIDE_H = 10.833333333333334


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.31, 0.24, 9, bold=True)


def add_picture(slide, filename: str, left: float, top: float, width: float, height: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide1, "Extended Data Figure 2 (continues)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture(slide1, "FigureS2A_KEGG_EWEvsESM_dotplot.png", 0.207, 0.468, 7.083, 2.944)
    add_picture(slide1, "FigureS2B_UpSetPlot_GO_BP_Terms.png", 0.207, 3.413, 7.083, 2.944)
    add_panel_label(slide1, "A", 0.208, 0.296)
    add_panel_label(slide1, "B", 0.208, 3.413)

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide2, "Extended Data Figure 2 (continued)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture(slide2, "FigureS2C_GO_BP_EWEvsESM_EWEvsAWE_dotplot.png", 0.207, 0.515, 3.542, 4.132)
    add_picture(slide2, "FigureS2D_KEGG_EWEvsESM_EWEvsAWE_dotplot.png", 3.743, 0.515, 3.542, 4.132)
    add_picture(slide2, "FigureS2E_RepresentativeGene_Boxplot.png", 0.203, 4.760, 7.083, 2.361)
    add_panel_label(slide2, "C", 0.208, 0.296)
    add_panel_label(slide2, "D", 3.736, 0.296)
    add_panel_label(slide2, "E", 0.208, 4.652)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
