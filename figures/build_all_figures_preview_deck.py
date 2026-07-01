#!/usr/bin/env python3
"""Build a preview deck containing previews of all figures."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


from path_config import BASE_DIR, FIG_DIR
OUT_PPTX = FIG_DIR / "AllFigures_preview_deck.pptx"

SLIDE_W = 10.0
SLIDE_H = 7.5


SLIDES = [
    ("Figure 1", ["Figure1_editable_preview.png"]),
    ("Figure 2", ["Figure2_editable_preview.png"]),
    ("Figure 3", ["Figure3_editable_slide1_preview.png", "Figure3_editable_slide2_preview.png"]),
    ("Figure 4", ["Figure4_editable_preview.png"]),
    ("Figure 5", ["Figure5_editable_slide1_preview.png", "Figure5_editable_slide2_preview.png"]),
    ("Figure 6", ["Figure6_editable_preview.png"]),
    ("Extended Data Figure 1", ["ExtendedDataFigure1_editable_slide1_preview.png", "ExtendedDataFigure1_editable_slide2_preview.png"]),
    ("Extended Data Figure 2", ["ExtendedDataFigure2_editable_slide1_preview.png", "ExtendedDataFigure2_editable_slide2_preview.png"]),
    ("Extended Data Figure 3", ["ExtendedDataFigure3_editable_slide1_preview.png", "ExtendedDataFigure3_editable_slide2_preview.png", "ExtendedDataFigure3_editable_slide3_preview.png"]),
    ("Extended Data Figure 4", ["ExtendedDataFigure4_editable_slide1_preview.png", "ExtendedDataFigure4_editable_slide2_preview.png"]),
]


def add_text(slide, text: str, left: float, top: float, width: float, height: float,
             size_pt: float, bold: bool = False) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_fit_picture(slide, image_path: Path, left: float, top: float, max_w: float, max_h: float) -> None:
    from PIL import Image

    with Image.open(image_path) as img:
        w_px, h_px = img.size
    aspect = w_px / h_px
    w = max_w
    h = w / aspect
    if h > max_h:
        h = max_h
        w = h * aspect
    x = left + (max_w - w) / 2
    y = top + (max_h - h) / 2
    slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    for title, files in SLIDES:
      slide = prs.slides.add_slide(prs.slide_layouts[6])
      add_text(slide, title, 0.15, 0.10, 5.0, 0.28, 11, bold=True)
      if len(files) == 1:
          add_fit_picture(slide, FIG_DIR / files[0], 0.25, 0.48, 9.5, 6.75)
      elif len(files) == 2:
          add_fit_picture(slide, FIG_DIR / files[0], 0.25, 0.55, 4.65, 6.55)
          add_fit_picture(slide, FIG_DIR / files[1], 5.10, 0.55, 4.65, 6.55)
      else:
          add_fit_picture(slide, FIG_DIR / files[0], 0.20, 0.55, 3.10, 6.55)
          add_fit_picture(slide, FIG_DIR / files[1], 3.45, 0.55, 3.10, 6.55)
          add_fit_picture(slide, FIG_DIR / files[2], 6.70, 0.55, 3.10, 6.55)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
