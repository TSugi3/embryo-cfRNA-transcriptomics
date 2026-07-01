#!/usr/bin/env python3
"""Render lightweight PNG previews from figure PPTX files.

This is a QA helper for the figure workflow. It renders embedded images and
native text boxes sufficiently for checking layout, clipping and legibility.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


EMU_PER_INCH = 914400


def emu_to_px(value: int, dpi: int) -> int:
    return int(round(value / EMU_PER_INCH * dpi))


def font_for(size_pt: float, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica.ttc",
        "/Library/Fonts/Arial.ttf",
    ]
    for path in candidates:
        if path and Path(path).exists():
            try:
                return ImageFont.truetype(path, max(6, int(round(size_pt * 2))))
            except OSError:
                continue
    return ImageFont.load_default()


def draw_text_shape(canvas: Image.Image, shape, dpi: int, scale: float) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    text_frame = shape.text_frame
    left = int(emu_to_px(shape.left, dpi) * scale)
    top = int(emu_to_px(shape.top, dpi) * scale)
    draw = ImageDraw.Draw(canvas)
    y = top
    for paragraph in text_frame.paragraphs:
        x = left
        if not paragraph.runs:
            y += 16
            continue
        for run in paragraph.runs:
            text = run.text or ""
            if not text:
                continue
            size_pt = float(run.font.size.pt) if run.font.size is not None else 10.0
            bold = bool(run.font.bold)
            font = font_for(size_pt, bold=bold)
            draw.text((x, y), text, fill=(0, 0, 0), font=font)
            bbox = draw.textbbox((x, y), text, font=font)
            x = bbox[2]
        y += 18


def render_pptx(pptx_path: Path, out_dir: Path, dpi: int = 192) -> list[Path]:
    prs = Presentation(str(pptx_path))
    scale = 1.0
    slide_w = emu_to_px(prs.slide_width, dpi)
    slide_h = emu_to_px(prs.slide_height, dpi)
    rendered: list[Path] = []
    for idx, slide in enumerate(prs.slides, start=1):
        canvas = Image.new("RGB", (slide_w, slide_h), "white")
        for shape in sorted(slide.shapes, key=lambda s: s.shape_id):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                blob = shape.image.blob
                with Image.open(__import__("io").BytesIO(blob)) as im:
                    im = im.convert("RGBA")
                    left = int(emu_to_px(shape.left, dpi) * scale)
                    top = int(emu_to_px(shape.top, dpi) * scale)
                    width = int(emu_to_px(shape.width, dpi) * scale)
                    height = int(emu_to_px(shape.height, dpi) * scale)
                    im = im.resize((width, height), Image.Resampling.LANCZOS)
                    canvas.paste(im.convert("RGB"), (left, top), im)
            elif getattr(shape, "has_text_frame", False):
                draw_text_shape(canvas, shape, dpi, scale)
        suffix = "" if len(prs.slides) == 1 else f"_slide{idx:02d}"
        out = out_dir / f"{pptx_path.stem}{suffix}_preview.png"
        canvas.save(out, dpi=(dpi, dpi))
        rendered.append(out)
    return rendered


def make_contact_sheet(images: list[Path], out_path: Path, thumb_w: int = 520) -> None:
    thumbs = []
    for path in images:
        with Image.open(path) as im:
            im = im.convert("RGB")
            h = int(im.height * thumb_w / im.width)
            im = im.resize((thumb_w, h), Image.Resampling.LANCZOS)
            thumbs.append((path, im))
    cols = 3
    label_h = 30
    gap = 22
    rows = (len(thumbs) + cols - 1) // cols
    row_h = max((im.height + label_h for _, im in thumbs), default=1)
    sheet = Image.new("RGB", (cols * thumb_w + (cols + 1) * gap, rows * row_h + (rows + 1) * gap), "white")
    draw = ImageDraw.Draw(sheet)
    font = font_for(9)
    for i, (path, im) in enumerate(thumbs):
        r, c = divmod(i, cols)
        x = gap + c * (thumb_w + gap)
        y = gap + r * (row_h + gap)
        draw.text((x, y), path.stem.replace("_preview", ""), fill=(0, 0, 0), font=font)
        sheet.paste(im, (x, y + label_h))
    sheet.save(out_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--figure-dir", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--contact-sheet", required=True, type=Path)
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    pptx_files = [
        "Figure1_editable.pptx",
        "Figure2_editable.pptx",
        "Figure3_editable.pptx",
        "Figure4_editable.pptx",
        "Figure5_editable.pptx",
        "Figure6_editable.pptx",
        "ExtendedDataFigure1_editable.pptx",
        "ExtendedDataFigure2_editable.pptx",
        "ExtendedDataFigure3_editable.pptx",
        "ExtendedDataFigure4_editable.pptx",
    ]
    all_images: list[Path] = []
    for name in pptx_files:
        all_images.extend(render_pptx(args.figure_dir / name, args.out_dir))
    make_contact_sheet(all_images, args.contact_sheet)
    print(f"Rendered {len(all_images)} preview images")
    print(args.contact_sheet)


if __name__ == "__main__":
    main()
