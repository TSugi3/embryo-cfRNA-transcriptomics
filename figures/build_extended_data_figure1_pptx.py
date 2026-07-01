#!/usr/bin/env python3
"""Build revised Extended Data Figure 1 as a two-slide working PowerPoint file."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


from path_config import BASE_DIR, FIG_DIR
OUT_PPTX = FIG_DIR / "ExtendedDataFigure1_Revised_working.pptx"

SLIDE_W = 7.5
SLIDE_H = 10.833333333333334


def add_textbox(slide, text: str, left: float, top: float, width: float, height: float,
                size_pt: float, bold: bool = False) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_panel_label(slide, label: str, left: float, top: float) -> None:
    add_textbox(slide, label, left, top, 0.31, 0.24, 9, bold=True)


def add_picture(slide, filename: str, left: float, top: float, width: float, height: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def add_picture_keep_aspect(slide, filename: str, left: float, top: float, width: float) -> None:
    path = FIG_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide1, "Extended Data Figure 1 (continues)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture(slide1, "FigureS1A_Barplot_ReadCounts.png", 0.208, 0.426, 7.083, 2.000)
    add_picture(slide1, "FigureS1B_ViolinPlot_RawReads.png", 0.207, 2.417, 2.347, 2.000)
    add_picture(slide1, "FigureS1C_MappingRate_PerGroup.png", 2.550, 2.417, 2.347, 2.000)
    add_picture(slide1, "FigureS1D_GiniCoefficient_withPval.png", 5.002, 2.417, 2.145, 2.000)
    add_picture_keep_aspect(slide1, "FigureS1E_Subsampling_Gene_Count_ALL_log10.png", 0.204, 4.383, 3.400)
    add_picture_keep_aspect(slide1, "FigureS1F_Subsampling_Gene_Count_ProteinCoding_log10.png", 3.892, 4.383, 3.400)
    for label, left, top in [
        ("A", 0.208, 0.296), ("B", 0.208, 2.383), ("C", 2.556, 2.383),
        ("D", 4.899, 2.383), ("E", 0.192, 4.348), ("F", 3.870, 4.348)
    ]:
        add_panel_label(slide1, label, left, top)

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide2, "Extended Data Figure 1 (continued)", 0.0, 0.005, 3.09, 0.30, 10)
    add_picture(slide2, "FigureS1G_MAplot_Facet.png", 0.208, 0.308, 3.542, 3.542)
    add_picture(slide2, "FigureS1H_DensityPlot_Facet.png", 3.750, 0.308, 3.542, 3.542)
    add_picture(slide2, "FigureS1I_UpSetPlot_DEG.png", 0.208, 3.847, 7.083, 3.139)
    add_panel_label(slide2, "G", 0.208, 0.296)
    add_panel_label(slide2, "H", 3.750, 0.308)
    add_panel_label(slide2, "I", 0.208, 3.893)

    prs.save(str(OUT_PPTX))
    print(OUT_PPTX)


if __name__ == "__main__":
    main()
